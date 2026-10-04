"""
tests/test_excel_logic.py

El Excel se arma leyendo el PDF ya generado (logic/excel_logic.py). Para que
estos tests fallen si el PDF cambia, el PDF de muestra se genera con el
`_tabla_vehiculo` REAL de logic/pdf_logic.py (db=None: sin BD). Puro: sin BD.
"""
import io
import os
import sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
import pytest
from reportlab.lib.pagesizes import LETTER, portrait
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Spacer

from logic import pdf_logic as P
from logic.excel_logic import leer_pdf, pdf_a_excel


# ── Fixtures ──────────────────────────────────────────────────
def _suc(orden, nombre, peso, num_tienda=None):
    return {"orden": orden, "nombre": nombre, "peso_kg": peso, "num_tienda": num_tienda}


def _may(orden, nombre, peso):
    return {"orden": orden, "nombre": nombre, "documento": nombre, "peso_kg": peso}


def _ruta(dia, sucursales, mayoristas=(), pct=50.0):
    peso = sum(s["peso_kg"] for s in sucursales) + sum(m["peso_kg"] for m in mayoristas)
    return {"dia": dia, "peso_kg": peso, "pct_utilizacion": pct, "tipo": "ruta",
            "sucursales": list(sucursales), "mayoristas": list(mayoristas)}


def _pdf(vehiculos, vol_map=None) -> bytes:
    """
    PDF con el mismo documento de dos columnas que arma generar_pdf
    (pdf_logic.py, 'doc_pdf'), pero alimentado con rutas en memoria.
    vehiculos: [(abrev, placas, chofer, ton, [rutas])]
    """
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=portrait(LETTER),
                          rightMargin=P.MARGEN, leftMargin=P.MARGEN,
                          topMargin=65, bottomMargin=P.MARGEN)
    marcos = [
        Frame(P.MARGEN + i * (P.ANCHO_COL + P.ESPACIO_ENTRE_COLS), P.MARGEN,
              P.ANCHO_COL, P.PH - 75,
              leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id=f"c{i}")
        for i in range(2)
    ]
    doc.addPageTemplates([PageTemplate(
        id="DosColumnas", frames=marcos,
        onPage=P._draw_header("Semana de prueba  2026-10-05 — 2026-10-09",
                              "04/10/2026 10:00:00"))])
    elementos = []
    for abrev, placas, chofer, ton, rutas in vehiculos:
        elementos.extend(P._tabla_vehiculo(abrev, placas, rutas, chofer, ton,
                                           None, vol_map or {}, None))
    doc.build(elementos)
    return buf.getvalue()


def _pdf_sin_tablas() -> bytes:
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=portrait(LETTER))
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(20, 20, 500, 700)])])
    doc.build([Spacer(1, 10)])
    return buf.getvalue()


def _veh_simple(abrev, chofer, n_paradas=4):
    sucs = [_suc(i, f"Sucursal {abrev} {i}", 100 * i) for i in range(1, n_paradas + 1)]
    return (abrev, "XX-000", chofer, 4.0, [_ruta("martes", sucs)])


# ── leer_pdf ──────────────────────────────────────────────────
def test_cabecera_se_lee_del_pdf():
    datos = leer_pdf(_pdf([_veh_simple("F 350_1", "CESAR")]))
    assert datos["cabecera"][0] == "INTEGRADORA COMERCIAL DEL GOLFO"
    assert "Semana de prueba" in datos["cabecera"][1]
    assert datos["cabecera"][2] == "Expedido: 04/10/2026 10:00:00"


def test_titulo_del_vehiculo_incluye_chofer_y_toneladas():
    datos = leer_pdf(_pdf([_veh_simple("F 350_1", "CESAR")]))
    assert datos["vehiculos"][0]["titulo"] == "F 350_1 · CESAR (4 ton)"


def test_paradas_con_peso_volumen_y_porcentaje():
    ruta = _ruta("martes",
                 [_suc(1, "Cosamaloapan 1", 861, num_tienda=11),
                  _suc(2, "Cosamaloapan 2", 545, num_tienda=12)],
                 pct=35.2)
    datos = leer_pdf(_pdf([("F 350_1", "XX", "CESAR", 4.0, [ruta])],
                          vol_map={11: 1.209, 12: 0.909}))
    dia = datos["vehiculos"][0]["dias"][0]
    assert [p["parada"] for p in dia["paradas"]] == ["Cosamaloapan 1", "Cosamaloapan 2"]
    assert [p["sec"] for p in dia["paradas"]] == ["1", "2"]
    assert [p["peso"] for p in dia["paradas"]] == [861.0, 545.0]
    assert [p["vol"] for p in dia["paradas"]] == [1.209, 0.909]
    assert dia["paradas"][0]["pct"] == pytest.approx(0.61)   # el PDF muestra 61% (sin decimales)


def test_fila_total_por_dia():
    ruta = _ruta("martes", [_suc(1, "A", 861, 11), _suc(2, "B", 545, 12)], pct=35.2)
    datos = leer_pdf(_pdf([("F 350_1", "XX", "CESAR", 4.0, [ruta])],
                          vol_map={11: 1.0, 12: 0.5}))
    total = datos["vehiculos"][0]["dias"][0]["total"]
    assert total["parada"].startswith("TOTAL MARTES")
    assert total["peso"] == 1406.0
    assert total["vol"] == 1.5
    assert total["pct"] == pytest.approx(0.352)


def test_volumen_ausente_se_lee_como_none():
    ruta = _ruta("martes", [_suc(1, "Sin volumen", 100)])
    datos = leer_pdf(_pdf([("F 350_1", "XX", "CESAR", 4.0, [ruta])]))
    assert datos["vehiculos"][0]["dias"][0]["paradas"][0]["vol"] is None


def test_mayorista_se_detecta_por_el_color_naranja():
    ruta = _ruta("lunes",
                 [_suc(1, "Piedras Negras", 951), _suc(3, "Tlalixcoyan", 405)],
                 [_may(2, "AA2111_CLAUDIA GUEVARA", 21)])
    datos = leer_pdf(_pdf([("J 19", "XX", "VICTOR", 2.5, [ruta])]))
    paradas = datos["vehiculos"][0]["dias"][0]["paradas"]
    assert [(p["parada"], p["mayorista"]) for p in paradas] == [
        ("Piedras Negras", False),
        ("AA2111_CLAUDIA GUEVARA", True),
        ("Tlalixcoyan", False),
    ]
    assert "1 may." in datos["vehiculos"][0]["dias"][0]["total"]["parada"]


def test_varios_dias_dia_y_apoyo_correctos_y_en_orden():
    martes = _ruta("martes", [_suc(i, f"M{i}", 400) for i in range(1, 6)])   # 5 suc -> 2 aux
    jueves = _ruta("jueves", [_suc(1, "J1", 300), _suc(2, "J2", 200)])       # 500 kg -> 1 aux
    datos = leer_pdf(_pdf([("F 350_1", "XX", "CESAR", 4.0, [jueves, martes])]))
    dias = datos["vehiculos"][0]["dias"]
    assert [(d["dia"], d["apoyo"], len(d["paradas"])) for d in dias] == [
        ("MARTES", "2", 5), ("JUEVES", "1", 2)]


def test_varios_vehiculos_en_dos_columnas_y_paginas_conservan_el_orden():
    vehiculos = [_veh_simple(f"V{i:02d}", f"CHOFER{i}", n_paradas=8) for i in range(1, 13)]
    datos = leer_pdf(_pdf(vehiculos))
    assert [v["titulo"].split()[0] for v in datos["vehiculos"]] == [f"V{i:02d}" for i in range(1, 13)]
    assert all(len(v["dias"][0]["paradas"]) == 8 for v in datos["vehiculos"])


def test_pdf_sin_tablas_de_rutas_levanta_valueerror():
    with pytest.raises(ValueError, match="no contiene tablas"):
        leer_pdf(_pdf_sin_tablas())


# ── pdf_a_excel ───────────────────────────────────────────────
def _hoja(bytes_pdf):
    return openpyxl.load_workbook(io.BytesIO(pdf_a_excel(bytes_pdf))).active


def _fila_con(ws, texto, col=4):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == texto:
            return r
    raise AssertionError(f"no hay fila con {texto!r} en la columna {col}")


def test_excel_tiene_una_hoja_resumen_con_cabecera():
    ws = _hoja(_pdf([_veh_simple("F 350_1", "CESAR")]))
    assert ws.title == "RESUMEN"
    assert ws["A1"].value == "INTEGRADORA COMERCIAL DEL GOLFO"
    assert ws["A3"].value == "Expedido: 04/10/2026 10:00:00"
    assert ws["A5"].value == "F 350_1 · CESAR (4 ton)"


def test_excel_numeros_son_numeros_y_volumen_ausente_queda_vacio():
    ruta = _ruta("martes", [_suc(1, "Con vol", 1286, 11), _suc(2, "Sin vol", 150)], pct=40.0)
    ws = _hoja(_pdf([("F 350_1", "XX", "CESAR", 4.0, [ruta])], vol_map={11: 1.338}))
    r1, r2 = _fila_con(ws, "Con vol"), _fila_con(ws, "Sin vol")
    assert ws.cell(r1, 5).value == 1286 and isinstance(ws.cell(r1, 5).value, int)
    assert ws.cell(r1, 6).value == pytest.approx(1.338)
    assert ws.cell(r1, 7).value == pytest.approx(0.90)   # 1286/1436 = 89.5%; el PDF lo muestra como 90%
    assert ws.cell(r2, 6).value is None


def test_excel_bloques_apilados_verticalmente_con_fila_en_blanco():
    ws = _hoja(_pdf([_veh_simple("F 350_1", "CESAR"), _veh_simple("J 19", "VICTOR")]))
    r1 = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("F 350_1"))
    r2 = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("J 19"))
    assert r2 > r1
    assert all(ws.cell(r2 - 1, c).value is None for c in range(1, 8))   # fila en blanco


def test_excel_combina_dia_y_apoyo_y_los_pone_en_la_primera_fila_del_dia():
    ws = _hoja(_pdf([_veh_simple("F 350_1", "CESAR", n_paradas=4)]))
    primera = _fila_con(ws, "Sucursal F 350_1 1")
    assert ws.cell(primera, 1).value == "MARTES"
    assert ws.cell(primera, 2).value == 1          # 1000 kg, 4 suc -> 1 auxiliar
    rangos = {str(r) for r in ws.merged_cells.ranges}
    assert f"A{primera}:A{primera + 3}" in rangos
    assert f"B{primera}:B{primera + 3}" in rangos


def test_excel_mayorista_naranja_y_total_con_fondo_naranja():
    ruta = _ruta("lunes", [_suc(1, "Normal", 500)], [_may(2, "AA1_MAYORISTA", 50)])
    ws = _hoja(_pdf([("J 19", "XX", "VICTOR", 2.5, [ruta])]))
    rm, rn = _fila_con(ws, "AA1_MAYORISTA"), _fila_con(ws, "Normal")
    assert ws.cell(rm, 4).font.color.rgb.endswith("EA580C")
    assert not ws.cell(rn, 4).font.color.rgb.endswith("EA580C")
    total = next(r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(r, 4).value or "").startswith("TOTAL"))
    assert ws.cell(total, 4).fill.start_color.rgb.endswith("FFEDD5")


def test_excel_total_sobre_100_por_ciento_en_rojo():
    ruta = _ruta("martes", [_suc(1, "Pesada", 5000)], pct=125.0)
    ws = _hoja(_pdf([("F 350_1", "XX", "CESAR", 4.0, [ruta])]))
    total = next(r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(r, 4).value or "").startswith("TOTAL"))
    assert ws.cell(total, 7).value == pytest.approx(1.25)
    assert ws.cell(total, 7).font.color.rgb.endswith("C0392B")
