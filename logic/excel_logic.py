"""
logic/excel_logic.py
Convierte el PDF de reporte de pesos (generado por logic/pdf_logic.py) en un
Excel vertical: un bloque por vehículo, apilados uno debajo del otro.

El Excel es una reinterpretación del PDF YA generado: este módulo no consulta
la BD ni depende de pdf_logic. Si cambia el diseño del PDF (columnas, colores),
este lector debe ajustarse; tests/test_excel_logic.py genera el PDF con el
código real de pdf_logic y falla si dejan de calzar.
"""
import io
import re

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ── Lectura del PDF ───────────────────────────────────────────
NCOLS = 7                      # DIA, APOYO, SEC., PARADA, PESO, VOL, % RUTA
I_DIA, I_APOYO, I_SEC, I_PARADA, I_PESO, I_VOL, I_PCT = range(NCOLS)
ALTO_CABECERA_PDF = 58         # pt desde arriba: _draw_header dibuja hasta y=58
COLOR_MAYORISTA = (0xEA / 255, 0x58 / 255, 0x0C / 255)   # C_MAY_TEXT de pdf_logic
TOL_COLOR = 0.06
_VACIOS = {"", "—", "–", "-"}


def _color_cerca(color, objetivo) -> bool:
    if not color or len(color) != 3:
        return False
    return all(abs(a - b) <= TOL_COLOR for a, b in zip(color, objetivo))


def _txt(celda) -> str:
    return (celda or "").replace("\n", " ").strip()


def _num(texto: str):
    """'1,286' -> 1286.0 ; '0.843' -> 0.843 ; '—' -> None."""
    t = _txt(texto).replace(",", "")
    if t in _VACIOS:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def _pct(texto: str):
    """'28%' -> 0.28 ; '80.6%' -> 0.806 ; '—' -> None."""
    t = _txt(texto).replace("%", "")
    n = _num(t)
    return None if n is None else n / 100


def _es_mayorista(chars, bbox) -> bool:
    """True si el texto de la celda `bbox` está en naranja (C_MAY_TEXT)."""
    x0, top, x1, bottom = bbox
    en_celda = [c for c in chars
                if c["x0"] >= x0 - 1 and c["x1"] <= x1 + 1
                and c["top"] >= top - 1 and c["bottom"] <= bottom + 1
                and c["text"].strip()]
    return bool(en_celda) and all(
        _color_cerca(c.get("non_stroking_color"), COLOR_MAYORISTA) for c in en_celda)


def _filas_de_tabla(page, tabla) -> list:
    """Filas de datos de una tabla de vehículo (sin título ni encabezado)."""
    chars = page.chars
    filas = []
    for fila, celdas in list(zip(tabla.extract(), tabla.rows))[2:]:
        bbox_parada = celdas.cells[I_PARADA]
        filas.append({
            "dia":       _txt(fila[I_DIA]),
            "apoyo":     _txt(fila[I_APOYO]),
            "sec":       _txt(fila[I_SEC]),
            "parada":    _txt(fila[I_PARADA]),
            "peso":      _num(fila[I_PESO]),
            "vol":       _num(fila[I_VOL]),
            "pct":       _pct(fila[I_PCT]),
            "mayorista": bool(bbox_parada) and _es_mayorista(chars, bbox_parada),
        })
    return filas


def _agrupar_dias(filas: list) -> list:
    """Parte las filas en días: cada día termina en su fila 'TOTAL ...'."""
    dias, grupo = [], []
    for f in filas:
        grupo.append(f)
        if f["parada"].upper().startswith("TOTAL"):
            paradas, total = grupo[:-1], grupo[-1]
            dia = next((g["dia"] for g in paradas if g["dia"]), "")
            if not dia:
                m = re.match(r"TOTAL\s+(\S+)", total["parada"], re.I)
                dia = m.group(1) if m else ""
            apoyo = next((g["apoyo"] for g in paradas if g["apoyo"]), "")
            dias.append({"dia": dia, "apoyo": apoyo, "paradas": paradas, "total": total})
            grupo = []
    return dias


def leer_pdf(bytes_pdf: bytes) -> dict:
    """
    -> {"cabecera": [líneas], "vehiculos": [{"titulo": str, "dias": [...]}]}
    Los vehículos salen en el orden de lectura del PDF (por página: columna
    izquierda de arriba abajo y luego la derecha).
    """
    with pdfplumber.open(io.BytesIO(bytes_pdf)) as pdf:
        if not pdf.pages:
            raise ValueError("El PDF no contiene tablas de rutas.")
        cab = pdf.pages[0].crop((0, 0, pdf.pages[0].width, ALTO_CABECERA_PDF))
        cabecera = [l.strip() for l in (cab.extract_text() or "").splitlines() if l.strip()]

        bloques = []   # [{"titulo", "filas"}], una entrada por vehículo
        for page in pdf.pages:
            tablas = [t for t in page.find_tables()
                      if len(t.extract()) >= 3 and len(t.extract()[1]) == NCOLS]
            tablas.sort(key=lambda t: (0 if t.bbox[0] < page.width / 2 else 1, t.bbox[1]))
            for t in tablas:
                titulo = _txt(t.extract()[0][0])
                filas = _filas_de_tabla(page, t)
                # Continuación: reportlab repite título+encabezado si una
                # tabla se parte entre columnas/páginas.
                if bloques and bloques[-1]["titulo"] == titulo:
                    bloques[-1]["filas"].extend(filas)
                else:
                    bloques.append({"titulo": titulo, "filas": filas})

    if not bloques:
        raise ValueError("El PDF no contiene tablas de rutas.")
    return {
        "cabecera": cabecera,
        "vehiculos": [{"titulo": b["titulo"], "dias": _agrupar_dias(b["filas"])}
                      for b in bloques],
    }


# ── Escritura del Excel ───────────────────────────────────────
ANCHOS = [11, 8, 7, 46, 11, 10, 10]
AZUL, AZUL_CLARO, AZUL_DIA, AZUL_TOTAL = "1565C0", "E3F2FD", "F5F9FF", "BBDEFB"
NARANJA, NARANJA_FONDO, NARANJA_TOTAL = "EA580C", "FFF7ED", "FFEDD5"
ROJO = "C0392B"
_LADO = Side(style="thin", color="B0BEC5")
BORDE = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)


def _fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def _entero_o_texto(valor: str):
    """APOYO '2' -> 2 ; '—' -> '—'."""
    return int(valor) if valor.isdigit() else valor


def _escribir_bloque(ws, fila: int, veh: dict) -> int:
    """Escribe un vehículo desde `fila`; devuelve la siguiente fila libre."""
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=NCOLS)
    c = ws.cell(fila, 1, veh["titulo"])
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = _fill(AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, NCOLS + 1):
        ws.cell(fila, col).border = BORDE
        ws.cell(fila, col).fill = _fill(AZUL)
    fila += 1

    for col, etiqueta in enumerate(
            ["DIA", "APOYO", "SEC.", "PARADA", "PESO (KG)", "VOL (m³)", "% RUTA"], 1):
        c = ws.cell(fila, col, etiqueta)
        c.font = Font(bold=True, color=AZUL, size=9)
        c.fill = _fill(AZUL_CLARO)
        c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fila += 1

    for dia in veh["dias"]:
        inicio = fila
        for p in dia["paradas"]:
            _escribir_parada(ws, fila, p)
            fila += 1
        if fila - 1 > inicio:
            ws.merge_cells(start_row=inicio, start_column=1, end_row=fila - 1, end_column=1)
            ws.merge_cells(start_row=inicio, start_column=2, end_row=fila - 1, end_column=2)
        ws.cell(inicio, 1, dia["dia"])
        ws.cell(inicio, 2, _entero_o_texto(dia["apoyo"]))
        for col in (1, 2):
            c = ws.cell(inicio, col)
            c.font = Font(bold=True, size=9)
            c.alignment = Alignment(horizontal="center", vertical="center")
        _escribir_total(ws, fila, dia["total"])
        fila += 1
    return fila


def _escribir_parada(ws, fila: int, p: dict) -> None:
    color = NARANJA if p["mayorista"] else "000000"
    for col in range(1, NCOLS + 1):
        c = ws.cell(fila, col)
        c.border = BORDE
        c.fill = _fill(NARANJA_FONDO if p["mayorista"] else
                       (AZUL_DIA if col <= I_APOYO + 1 else "FFFFFF"))
    ws.cell(fila, I_SEC + 1, _entero_o_texto(p["sec"]))
    ws.cell(fila, I_PARADA + 1, p["parada"])
    _numero(ws, fila, I_PESO + 1, p["peso"], "#,##0")
    _numero(ws, fila, I_VOL + 1, p["vol"], "0.000")
    _numero(ws, fila, I_PCT + 1, p["pct"], "0%")
    for col in range(I_SEC + 1, NCOLS + 1):
        c = ws.cell(fila, col)
        c.font = Font(size=9, color=color, bold=p["mayorista"])
        c.alignment = Alignment(
            horizontal="left" if col == I_PARADA + 1 else
                       "center" if col in (I_SEC + 1, I_PCT + 1) else "right",
            vertical="center")


def _escribir_total(ws, fila: int, t: dict) -> None:
    con_may = "may." in t["parada"]
    fondo = NARANJA_TOTAL if con_may else AZUL_TOTAL
    for col in range(1, NCOLS + 1):
        c = ws.cell(fila, col)
        c.fill = _fill(fondo)
        c.border = BORDE
        c.font = Font(bold=True, size=9)
    ws.cell(fila, I_PARADA + 1, t["parada"]).alignment = Alignment(horizontal="left")
    _numero(ws, fila, I_PESO + 1, t["peso"], "#,##0")
    _numero(ws, fila, I_VOL + 1, t["vol"], "0.000")
    _numero(ws, fila, I_PCT + 1, t["pct"], "0.0%")
    for col in (I_PESO + 1, I_VOL + 1):
        ws.cell(fila, col).alignment = Alignment(horizontal="right")
    pct = ws.cell(fila, I_PCT + 1)
    pct.alignment = Alignment(horizontal="center")
    pct.font = Font(bold=True, size=9,
                    color=ROJO if (t["pct"] or 0) > 1 else AZUL)


def _numero(ws, fila: int, col: int, valor, formato: str) -> None:
    c = ws.cell(fila, col, None if valor is None else
                (int(valor) if formato == "#,##0" else valor))
    c.number_format = formato


def pdf_a_excel(bytes_pdf: bytes) -> bytes:
    """PDF de reporte de pesos -> bytes de un .xlsx vertical (hoja RESUMEN)."""
    datos = leer_pdf(bytes_pdf)

    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    for i, ancho in enumerate(ANCHOS, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    fila = 1
    for i, linea in enumerate(datos["cabecera"]):
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=NCOLS)
        c = ws.cell(fila, 1, linea)
        c.alignment = Alignment(horizontal="center")
        c.font = (Font(bold=True, size=14) if i == 0 else
                  Font(color="808080", size=9) if linea.startswith("Expedido") else
                  Font(size=10))
        fila += 1
    fila += 1

    for veh in datos["vehiculos"]:
        fila = _escribir_bloque(ws, fila, veh) + 1   # una fila en blanco entre bloques

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
